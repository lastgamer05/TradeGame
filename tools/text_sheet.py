"""게임 안의 한국어 글(대사, 선택지, 이벤트, 컷신 자막, 안내, 인물·구역 이름)을 엑셀 시트로 뽑고,
시트에서 고친 문장을 다시 데이터에 넣는다.

    python tools/text_sheet.py export [대사시트.xlsx]    # 시트 만들기 (기본: 대사시트.xlsx)
    python tools/text_sheet.py import 대사시트.xlsx      # "고칠 문장" 칸을 채운 줄만 반영

- 시트는 대화 / 이벤트 / 연출·안내 / 이름·설명 네 장이다.
- "고칠 문장" 칸만 채우면 된다. 비워 둔 줄은 그대로 둔다. 줄바꿈은 셀 안 줄바꿈(Alt+Enter)으로 쓴다.
- {player}는 주인공 이름으로 바뀌는 자리라 지우지 않는다. 선택지 앞의 [교섭 DC 15] 같은 태그는 게임이 붙이므로 쓰지 않는다.
- "키" 열(숨김)은 데이터 위치라 건드리지 않는다.
- 반영할 때 JSON 파일의 서식은 그대로 두고 그 문장만 바꾼다. 시트를 뽑은 뒤 원문이 바뀐 줄은 건너뛰고 알려 준다.
- 반영한 뒤에는 테스트(CLAUDE.md의 명령)를 돌려 데이터 검증을 확인한다.
"""
import glob
import json
import os
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "TradeGame", "data")
BS = chr(92)

SHEETS = [
    ("대화", ["dialogues/arrival_*.json", "dialogues/*.json"]),
    ("이벤트", ["events/*.json"]),
    ("연출·안내", ["cutscenes/*.json", "guide.json"]),
    ("이름·설명", ["npcs.json", "locations.json", "quests.json", "companions.json", "politics.json", "cities.json",
                 "factions.json", "goods.json"]),
]
HEADERS = ["키", "파일", "위치", "화자·종류", "현재 문장", "고칠 문장"]


def hangul(s):
    return any("가" <= c <= "힣" for c in s)


def walk(o, p, out):
    """한글이 든 문자열을 (경로, 문장) 순서대로 모은다. 파일 안 등장 순서와 같다."""
    if isinstance(o, dict):
        for k, v in o.items():
            walk(v, p + [k], out)
    elif isinstance(o, list):
        for i, v in enumerate(o):
            walk(v, p + [i], out)
    elif isinstance(o, str) and hangul(o):
        out.append((p, o))


def rel(path):
    return os.path.relpath(path, ROOT).replace(BS, "/")


def files():
    seen = set()
    for name, patterns in SHEETS:
        lst = []
        for pat in patterns:
            for f in sorted(glob.glob(os.path.join(ROOT, pat))):
                r = rel(f)
                if r not in seen:
                    seen.add(r)
                    lst.append(r)
        yield name, lst


def load_names():
    names = {"player": "주인공", "narrator": "내레이션"}
    for npc in json.load(open(os.path.join(ROOT, "npcs.json"), encoding="utf8")):
        names[npc["id"]] = npc["name"]
    return names


def describe(data, p, names):
    """위치와 화자·종류 칸에 쓸 말."""
    keys = [str(k) for k in p]
    last = keys[-1]
    where = "/".join(keys[:-1])
    if keys[0] == "nodes" and len(keys) >= 3:
        node = data["nodes"][keys[1]]
        if "choices" in keys:
            if "fail" in last or "fail" in where:
                return keys[1], "선택지 결과"
            return keys[1], "선택지"
        return keys[1], names.get(node.get("speaker", ""), node.get("speaker", ""))
    if "choices" in keys:
        if "outcomes" in keys:
            return where, "결과 (%s)" % {"success": "성공", "failure": "실패", "partial": "부분 성공"}.get(keys[-2], keys[-2])
        return where, "선택지"
    if "lines" in keys:
        return where, "자막"
    return where, {"text": "본문", "title": "제목", "name": "이름", "description": "설명", "role": "역할",
                   "summary": "요약", "hint": "설명", "where": "장소", "subtitle": "부제", "leave_text": "떠날 때 대사",
                   "start_text": "시작 알림", "end_text": "끝 알림"}.get(last, last)


def export(out_path):
    names = load_names()
    wb = Workbook()
    wb.remove(wb.active)
    total = 0
    for sheet_name, flist in files():
        ws = wb.create_sheet(sheet_name)
        ws.append(HEADERS)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="DDD6C8")
        for r in flist:
            data = json.load(open(os.path.join(ROOT, r), encoding="utf8"))
            items = []
            walk(data, [], items)
            for p, text in items:
                where, kind = describe(data, p, names)
                ws.append([r + "|" + "/".join(map(str, p)), r, where, kind, text, ""])
                total += 1
        ws.column_dimensions["A"].hidden = True
        for col, w in zip("BCDEF", [30, 26, 16, 70, 70]):
            ws.column_dimensions[col].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row[4:6]:
                c.alignment = Alignment(wrap_text=True, vertical="top")
            row[5].fill = PatternFill("solid", fgColor="FFF6D8")
        ws.freeze_panes = "E2"
        ws.auto_filter.ref = ws.dimensions
    wb.save(out_path)
    print("시트 저장: %s (%d줄)" % (out_path, total))


def import_(in_path):
    wb = load_workbook(in_path)
    changes = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue
            key, _, _, _, old, new = (list(row) + [None] * 6)[:6]
            if new is None or str(new).strip() == "":
                continue
            new = str(new).replace("\r\n", "\n")
            if new == old:
                continue
            r, path = key.split("|", 1)
            changes.setdefault(r, {})[path] = (old, new)
    applied = 0
    skipped = []
    for r, edits in changes.items():
        f = os.path.join(ROOT, r)
        raw = open(f, encoding="utf8", newline="").read()
        items = []
        walk(json.loads(raw), [], items)
        count = {}
        spans = []
        for p, text in items:
            k = count.get(text, 0)
            count[text] = k + 1
            path = "/".join(map(str, p))
            if path not in edits:
                continue
            old, new = edits.pop(path)
            if old != text:
                skipped.append("%s %s: 시트를 뽑은 뒤 원문이 바뀌었다" % (r, path))
                continue
            tok = json.dumps(text, ensure_ascii=False)
            pos = -1
            for _ in range(k + 1):
                pos = raw.index(tok, pos + 1)
            spans.append((pos, pos + len(tok), json.dumps(new, ensure_ascii=False)))
        for path in edits:
            skipped.append("%s %s: 데이터에 이 자리가 없다" % (r, path))
        for a, b, rep in sorted(spans, reverse=True):
            raw = raw[:a] + rep + raw[b:]
        json.loads(raw)
        open(f, "w", encoding="utf8", newline="").write(raw)
        applied += len(spans)
    print("반영: %d줄" % applied)
    for s in skipped:
        print("건너뜀: " + s)


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    if len(sys.argv) < 2 or sys.argv[1] not in ("export", "import"):
        print(__doc__)
        sys.exit(1)
    if sys.argv[1] == "export":
        export(sys.argv[2] if len(sys.argv) > 2 else "대사시트.xlsx")
    else:
        import_(sys.argv[2])
